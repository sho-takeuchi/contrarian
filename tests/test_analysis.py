import copy
from datetime import date, datetime, timedelta
from decimal import Decimal as D
import json
import os
from pathlib import Path
import plistlib
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree

try:
    import sqlalchemy
    from pydantic import ValidationError
except ModuleNotFoundError:
    raise unittest.SkipTest('Install .[analysis] to run optional analysis tests.')

from jp_screening.core import ScreeningError, JST
from jp_screening.analysis.config import Config
from jp_screening.analysis.domain import features, outcome, summarize, quartiles, build_episodes, representative
from jp_screening.analysis.ingest import prepare_run
from jp_screening.analysis.rag import Intent, extract, split_document, Ollama
from jp_screening.analysis.schedule import mac_definition, windows_definition, manage
from jp_screening.analysis.tracking import latest_closed_day
from analysis_support import START, DAYS, CAL, NOW, quotes, run_folder, result


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.prices = {date.fromisoformat(x["Date"]):x for x in quotes()}

    def evaluate(self, **kwargs):
        args = dict(prices=self.prices, sessions=DAYS, start=START, horizon=5, basis="close", cutoff=DAYS[5])
        return outcome(**(args | kwargs))

    def test_business_day_horizon_skips_holiday_and_weekend(self):
        value = self.evaluate()
        self.assertEqual(value["target_date"], "2026-01-13")
        self.assertEqual(value["return_pct"], D(5))

    def test_both_baselines_and_excursions_exclude_signal_day(self):
        self.prices[START]["H"] = 1000
        self.prices[DAYS[1]]["O"] = 105
        value = self.evaluate(basis="next_open")
        self.assertEqual(value["return_pct"], 0)
        self.assertEqual(value["baseline"], 105)
        self.assertEqual(self.evaluate()["max_up_pct"], D(10))
        self.assertEqual(self.evaluate()["max_down_pct"], D(-10))

    def test_split_and_reverse_split(self):
        for event in (D("0.5"), D(2)):
            with self.subTest(event=event):
                prices = {date.fromisoformat(x["Date"]):x for x in quotes()}
                for i, day in enumerate(DAYS[:6]):
                    prices[day].update(O=D(100), H=D(110), L=D(90), C=D(100))
                    if i >= 3:
                        for field in ("O", "H", "L", "C"):
                            prices[day][field] *= event
                prices[DAYS[3]]["AdjFactor"] = event
                prices[DAYS[5]]["C"] = D(110)*event
                value = self.evaluate(prices=prices)
                self.assertEqual(value["return_pct"], D(10))
                self.assertEqual(value["max_up_pct"], D(10))
                self.assertEqual(value["max_down_pct"], D(-10))

    def test_signal_day_factor_does_not_adjust_itself(self):
        self.prices[START]["AdjFactor"] = D("0.01")
        self.assertEqual(self.evaluate()["return_pct"], D(5))

    def test_future_split_does_not_change_finished_horizon(self):
        self.prices[DAYS[6]] = dict(O=1, H=1, L=1, C=1, AdjFactor=D("0.01"))
        self.assertEqual(self.evaluate()["return_pct"], D(5))

    def test_pending_is_not_missing_or_zero(self):
        value = self.evaluate(cutoff=DAYS[4])
        self.assertEqual(value["status"], "評価待ち")
        self.assertIsNone(value["return_pct"])

    def test_missing_next_open_only_invalidates_that_baseline(self):
        self.prices[DAYS[1]]["O"] = None
        self.assertEqual(self.evaluate()["status"], "評価済み")
        self.assertEqual(self.evaluate(basis="next_open")["status"], "必要価格の欠損")

    def test_missing_record_never_forward_filled(self):
        del self.prices[DAYS[2]]
        self.assertEqual(self.evaluate()["status"], "必要価格の欠損")

    def test_missing_factor_cannot_be_assumed_one(self):
        self.prices[DAYS[2]]["AdjFactor"] = None
        self.assertEqual(self.evaluate()["status"], "調整係数の欠損")

    def test_null_quote_is_not_assumed_delisted(self):
        self.prices[DAYS[2]].update(O=None, H=None, L=None, C=None)
        value = self.evaluate()
        self.assertEqual(value["status"], "評価済み")
        self.assertEqual(value["excursion_status"], "必要価格の欠損")

    def test_latest_uses_actual_last_quote_not_current_day(self):
        self.assertEqual(self.evaluate(horizon=0, cutoff=DAYS[8])["target_date"], str(DAYS[5]))

    def test_latest_on_signal_day_is_zero_for_close_only(self):
        self.assertEqual(self.evaluate(horizon=0, cutoff=START)["return_pct"], 0)
        self.assertIsNone(self.evaluate(horizon=0, cutoff=START, basis="next_open")["return_pct"])

    def test_invalid_period_or_price_rejected(self):
        with self.assertRaises(ScreeningError):
            self.evaluate(horizon=7)
        self.prices[START]["C"] = 0
        with self.assertRaises(ScreeningError):
            self.evaluate()

    def test_feature_formula_and_missing_denominator(self):
        values = features(result())
        self.assertEqual(values["long_low_rise"], 25)
        self.assertEqual(values["ceiling_room"], 20)
        for denominator in (None, 0, -1):
            self.assertIsNone(features(result() | {"sma5_day":denominator})["daily_ma_gap"])

    def test_summary_and_tied_quartiles(self):
        rows = [dict(code="10000", status=status, return_pct=value, features={"daily_ma_gap":D(-2)})
                for status, value in [("評価済み",D(10)),("評価済み",D(0)),("評価待ち",None),("必要価格の欠損",None)]]
        value = summarize(rows)
        self.assertEqual((value["total"],value["valid"],value["pending"],value["missing"],value["unique_codes"]), (4,2,1,1,1))
        self.assertEqual(value["positive_pct"], 50)
        self.assertEqual(value["mean"], 5)
        self.assertEqual([r["total"] for r in quartiles(rows,"daily_ma_gap")["groups"]], [4,0,0,0])


class EventTests(unittest.TestCase):
    def events(self, states, indices=None):
        runs, rows = [], []
        for i, flags in enumerate(states):
            day = DAYS[indices[i] if indices else i]
            runs.append(dict(id=str(i), rules_key="rule", asof=day, stage=3, eligible=True, created_at=str(day)))
            if flags:
                rows.append(dict(run_id=str(i), code="10000", s1="通過", s2="通過", s3=flags))
        return [x for x in build_episodes(runs, rows, DAYS) if x["stage"] == 3]

    def test_continuous_pass_is_one_episode(self):
        self.assertEqual(len(self.events(["通過"]*3)), 1)

    def test_confirmed_failure_then_pass_starts_new_episode(self):
        self.assertEqual(len(self.events(["通過","不通過","通過"])), 2)

    def test_unknown_absent_security_and_unobserved_day_do_not_reset(self):
        for states, indices in [(["通過","判定不能","通過"],None),(["通過",None,"通過"],None),(["通過","通過"],[0,3])]:
            values = self.events(states,indices)
            self.assertEqual(len(values),1)
            self.assertTrue(values[0]["continuity_unknown"])

    def test_representative_ignores_sample_and_prefers_depth_then_latest(self):
        common = dict(rules_key="r", asof=START, eligible=True, stage=3)
        runs = [common | dict(id="a",created_at="18"),common | dict(id="b",created_at="19"),
                common | dict(id="c",created_at="20",stage=1),common | dict(id="d",created_at="21",eligible=False)]
        self.assertEqual(representative(runs)[0]["id"], "b")


class InputAndSchedulerTests(unittest.TestCase):
    def test_hash_tampering_and_candidate_disagreement(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = run_folder(tmp)
            (folder/"candidates.csv").write_text("altered",encoding="utf-8")
            with self.assertRaisesRegex(ScreeningError,"ハッシュ"):
                prepare_run(folder)
            legacy = run_folder(tmp,"legacy",legacy=True)
            path = legacy/"candidates.csv"
            path.write_text(path.read_text(encoding="utf-8-sig").splitlines()[0]+"\n",encoding="utf-8-sig")
            with self.assertRaisesRegex(ScreeningError,"候補CSV"):
                prepare_run(legacy)

    def test_unconfirmed_legacy_and_sample_not_eligible(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name, kw in [("legacy",dict(legacy=True)),("sample",dict(scope="sample",mode="sample")),
                             ("recompute",dict(mode="recompute")),("specified",dict(scope="specified"))]:
                self.assertFalse(prepare_run(run_folder(tmp,name,**kw))[0]["eligible"])

    def test_intent_cannot_contain_sql_or_unbounded_limits(self):
        for data in [dict(kind="sql"),dict(kind="history",sql="DROP TABLE"),dict(kind="history",limit=501),
                     dict(kind="performance",horizon=7),dict(kind="history",codes=["not_a_code"]),
                     dict(kind="history",start="2026-02-01",end="2026-01-01")]:
            with self.assertRaises((ValidationError, ScreeningError)):
                Intent(**data)

    def test_remote_ollama_configuration_rejected(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ,{"OLLAMA_BASE_URL":"https://example.com"}):
            with self.assertRaises(ScreeningError):
                Config.load(tmp)

    def test_mac_and_windows_schedule_is_15_minutes_and_user_only(self):
        root = Path("/tmp/example project")
        mac = mac_definition(root)
        self.assertEqual(mac["StartInterval"],900)
        self.assertIn(str(root),mac["ProgramArguments"])
        xml = ElementTree.fromstring(windows_definition(root))
        ns = {"t":"http://schemas.microsoft.com/windows/2004/02/mit/task"}
        self.assertEqual(xml.find(".//t:Interval",ns).text,"PT15M")
        self.assertEqual(xml.find(".//t:LogonType",ns).text,"InteractiveToken")
        self.assertEqual(xml.find(".//t:RunLevel",ns).text,"LeastPrivilege")

    def test_scheduler_dry_run_never_registers(self):
        with patch("jp_screening.analysis.schedule.sys.platform","darwin"), patch("subprocess.run") as run:
            definition = manage(Path.cwd(),"install",True)
            self.assertEqual(plistlib.loads(definition.encode())["StartInterval"],900)
            run.assert_not_called()

    def test_market_close_and_holiday(self):
        self.assertEqual(latest_closed_day(CAL,NOW.replace(hour=10)),DAYS[4])
        weekend = datetime(2026,1,11,19,tzinfo=JST)
        self.assertEqual(latest_closed_day(CAL,weekend),date(2026,1,9))

    def test_docx_paragraphs_tables_and_headings(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/"rules.docx"
            doc = Document()
            doc.add_heading("判定仕様",1)
            doc.add_paragraph("月足の条件です。")
            doc.add_table(1,2).rows[0].cells[0].text = "週足"
            doc.save(path)
            content = extract(path)
            self.assertIn("週足",content)
            self.assertEqual(split_document(content)[0][0],"判定仕様")

    def test_missing_optional_config_does_not_throw_from_screener(self):
        from jp_screening.analysis.integration import optional_import
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root/"run.json").write_text("{}")
            with patch.dict(os.environ,{"ANALYSIS_AUTO_IMPORT":"true"}), patch("jp_screening.analysis.config.Config.load",side_effect=ScreeningError("設定不正")):
                self.assertFalse(optional_import(root,root))
            with patch.dict(os.environ,{"ANALYSIS_AUTO_IMPORT":"false","OLLAMA_BASE_URL":"https://example.com"}):
                self.assertTrue(optional_import(root,root))
